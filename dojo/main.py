import sys
import re
from os.path import exists
from pptx import Presentation as PptxPresentation
from pptx.util import Inches


class Presentation:
    def __init__(self, raw_markdown: str):
        self.raw_markdown = raw_markdown
        self.slides = []
        self.global_template = None  # placeholder for later
        self.parse()

    def parse(self):
        raw_slides = [s.strip() for s in self.raw_markdown.split('---') if s.strip()]
        for i, raw_slide in enumerate(raw_slides, start=1):
            slide = Slide(raw_slide, slide_number=i)
            self.slides.append(slide)

    def export_to_pptx(self, output_path):
        pptx = PptxPresentation()

        for slide_obj in self.slides:
            if slide_obj.slide_number == 1:
                ppt_slide = pptx.slides.add_slide(pptx.slide_layouts[0])  # Title Slide
                ppt_slide.shapes.title.text = slide_obj.title
                if slide_obj.subtitle:
                    ppt_slide.placeholders[1].text = slide_obj.subtitle
            else:
                ppt_slide = pptx.slides.add_slide(pptx.slide_layouts[1])  # Content Slide
                ppt_slide.shapes.title.text = slide_obj.title or f"Slide {slide_obj.slide_number}"
                text_frame = ppt_slide.placeholders[1].text_frame
                text_frame.clear()

                for obj in slide_obj.objects:
                    obj.export(ppt_slide, text_frame)

        pptx.save(output_path)
        print(f"✔ Presentation saved to {output_path}")


class Slide:
    def __init__(self, raw_content: str, slide_number: int):
        self.raw_content = raw_content
        self.slide_number = slide_number
        self.title = ''
        self.subtitle = ''
        self.layout = None  # e.g., ColumnLayout
        self.objects = []
        self.parse()

    def parse(self):
        lines = self.raw_content.split('\n')
        for line in lines:
            if line.startswith('# '):
                self.title = line[2:].strip()
            elif line.startswith('## '):
                self.subtitle = line[3:].strip()
            elif line.startswith('||L ') or line.startswith('||R '):
                if not self.layout:
                    self.layout = ColumnLayout()
                self.layout.parse_line(line)
            elif line.startswith('#IMAGE('):
                image = ImageObject(line)
                self.objects.append(image)
            elif line.strip().startswith('-') or line.strip().startswith('*'):
                if not self.objects or not isinstance(self.objects[-1], BulletedList):
                    self.objects.append(BulletedList())
                self.objects[-1].add_line(line)
            else:
                if line.strip():
                    self.objects.append(TextBox(line))


class SlideObject:
    def export(self, ppt_slide, text_frame):
        pass  # base method


class BulletedList(SlideObject):
    def __init__(self):
        self.items = []

    def add_line(self, line):
        match = re.match(r'^(\s*)[-*] (.+)', line)
        if match:
            indent_str = match.group(1)
            indent = len(indent_str)
            level = indent // 2
            bullet_text = match.group(2).strip()
            self.items.append((level, bullet_text))

    def export(self, ppt_slide, text_frame):
        for level, text in self.items:
            p = text_frame.add_paragraph()
            p.text = text
            p.level = level


class ImageObject(SlideObject):
    def __init__(self, line):
        self.path, self.left, self.top, self.width, self.height = self.parse_line(line)

    def parse_line(self, line):
        print(f"Parsing image line: {line}")

        path_match = re.search(r'#IMAGE\(([^)]+)\)', line)
        if not path_match:
            raise ValueError(f"Invalid image line (missing #IMAGE(...)): {line}")

        path = path_match.group(1).strip()
        print(f"→ Parsed path: {path}")

        left, top, width, height = 1.0, 2.0, None, 3.0  # defaults

        bracket_match = re.search(r'\[([^\]]+)\]', line)
        if bracket_match:
            values_str = bracket_match.group(1)
            values = [v.strip() for v in values_str.split(',')]
            print(f"→ Parsed values: {values}")

            if len(values) >= 1:
                left = float(values[0])
            if len(values) >= 2:
                top = float(values[1])
            if len(values) >= 3:
                width = float(values[2])
            if len(values) >= 4:
                height = float(values[3])

        print(f"→ Final values: left={left}, top={top}, width={width}, height={height}")
        return path, left, top, width, height

    def export(self, ppt_slide, text_frame):
        try:
            ppt_slide.shapes.add_picture(
                self.path,
                Inches(self.left),
                Inches(self.top),
                width=Inches(self.width) if self.width else None,
                height=Inches(self.height)
            )
            print(f"✔ Added image: {self.path}")
        except Exception as e:
            print(f"⚠ Error adding image {self.path}: {e}")


class TextBox(SlideObject):
    def __init__(self, text):
        self.text = text.strip()

    def export(self, ppt_slide, text_frame):
        p = text_frame.add_paragraph()
        p.text = self.text


class ColumnLayout:
    def __init__(self):
        self.left = []
        self.right = []

    def parse_line(self, line):
        if line.startswith('||L '):
            self.left.append(line[4:].strip())
        elif line.startswith('||R '):
            self.right.append(line[4:].strip())


if __name__ == "__main__":
    input_file = "file.pmd"
    output_file = "output.pptx"

    with open(input_file, 'r') as f:
        raw_markdown = f.read()

    presentation = Presentation(raw_markdown)
    presentation.export_to_pptx(output_file)
