from pptx import Presentation
from pptx.util import Inches
import re

# Add logging
import logging
logging.basicConfig(filename='presentation_parser.log', level=logging.DEBUG)


class PresentationParser:
    def __init__(self):
        self.prs = Presentation()
        self.current_slide = None

    def parse_file(self, filename):
        with open(filename, 'r', encoding='utf-8') as file:
            content = file.read()
        # Log content for debugging
        logging.info(f'Parsed content: {content}')

        slides = [s.strip() for s in content.split('---') if s.strip()]
        logging.info(f'Found {len(slides)} slides')
        for slide in slides:
            logging.info(f'Processing slide:\n{slide}')

        # Process title slide
        self._process_title_slide(slides[0])

        # Process other slides
        for slide_content in slides[1:]:
            self._process_slide(slide_content)

    def _process_title_slide(self, content):
        lines = content.strip().split('\n')
        title, subtitle = '', ''

        for line in lines:
            if line.startswith('# '):
                title = line[2:].strip()
            elif line.startswith('## '):
                subtitle = line[3:].strip()

        slide_layout = self.prs.slide_layouts[0]
        slide = self.prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = title or 'Title'
        if subtitle:
            slide.placeholders[1].text = subtitle

    def _process_slide(self, content):
        lines = content.strip().split('\n')
        title = ''
        body_lines = []

        if lines and lines[0].startswith('# '):
            title = lines[0][2:].strip()
            body_lines = lines[1:]
        else:
            body_lines = lines

        slide_layout = self.prs.slide_layouts[1]
        slide = self.prs.slides.add_slide(slide_layout)
        self.current_slide = slide
        slide.shapes.title.text = title or 'Slide'

        left_column, right_column, normal_text = [], [], []

        for line in body_lines:
            line = line.strip()
            if not line:
                continue

            # Check for left/right column markers
            if line.startswith('||L '):
                left_column.append(line[4:].strip())
                continue
            elif line.startswith('||R '):
                right_column.append(line[4:].strip())
                continue

            # Check for images
            if line.startswith('!['):
                self._add_image(line)
                continue  # Skip adding to text

            # Otherwise, collect as normal text
            normal_text.append(line)

        # Fill main text box (normal text)
        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()
        for line in normal_text:
            if line.startswith('- '):
                p = text_frame.add_paragraph()
                p.text = line[2:].strip()
                p.level = 0
            else:
                p = text_frame.add_paragraph()
                p.text = line
                p.level = 0

        # Fill left/right column placeholders
        for shape in slide.placeholders:
            if shape.placeholder_format.idx == 1 and left_column:
                shape.text = '\n'.join(left_column)
            elif shape.placeholder_format.idx == 2 and right_column:
                shape.text = '\n'.join(right_column)

    def _add_image(self, line):
        match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if match:
            img_path = match.group(2)
            try:
                self.current_slide.shapes.add_picture(
                    img_path,
                    Inches(1),
                    Inches(2),
                    width=Inches(4)
                )
            except Exception as e:
                print(f"Error adding image {img_path}: {e}")

    def save(self, filename):
        self.prs.save(filename)


def main():
    parser = PresentationParser()
    parser.parse_file('file.pmd')
    parser.save('output.pptx')


if __name__ == '__main__':
    main()
